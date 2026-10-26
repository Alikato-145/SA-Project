"""Build the Haris SA Word report and sequence figures from the reviewed Markdown."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).parent
SOURCE = ROOT / "README.md"
FIGURES = ROOT / "figures"
OUTPUT = ROOT / "Haris-SA-Report.docx"
FONT_PATH = "/System/Library/Fonts/Supplemental/Tahoma.ttf"


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_PATH, size)


def short(value: str, limit: int = 48) -> str:
    return value if len(value) <= limit else value[: limit - 1] + "…"


def wrap_label(draw: ImageDraw.ImageDraw, value: str, label_font: ImageFont.FreeTypeFont,
               max_width: int) -> list[str]:
    lines = [""]
    for char in value:
        candidate = lines[-1] + char
        if lines[-1] and draw.textlength(candidate, font=label_font) > max_width:
            split = lines[-1].rfind(" ")
            if split > len(lines[-1]) // 3:
                rest = lines[-1][split + 1:] + char
                lines[-1] = lines[-1][:split].rstrip()
                lines.append(rest.lstrip())
            else:
                lines.append(char)
        else:
            lines[-1] = candidate
    return lines


def draw_sequence(number: str, title: str, block: list[str], output_dir: Path | None = None) -> Path | None:
    lanes: list[tuple[str, str]] = []
    messages: list[tuple[str, str, str, bool]] = []
    for raw in block:
        line = raw.strip()
        participant = re.match(r"(?:actor|participant)\s+(\w+)\s+as\s+(.+)", line)
        if participant:
            lanes.append((participant.group(1), participant.group(2)))
            continue
        message = re.match(r"(\w+)(-->>|->>)(\w+):\s*(.+)", line)
        if message:
            messages.append((message.group(1), message.group(3), message.group(4), message.group(2) == "-->>"))
    if len(lanes) < 2 or not messages:
        return None

    scale = 2
    focused = len(lanes) >= 5
    width = 1500 if focused else max(1450, 360 * len(lanes))
    message_size = 23 if focused else 15
    measurer = ImageDraw.Draw(Image.new("RGB", (1, 1)))
    labels = [wrap_label(measurer, label, font(message_size * scale), 410 * scale)
              for _, _, label, _ in messages]
    row_heights = [max(110 if focused else 125, 34 * len(lines) + 40)
                   for lines in labels]
    height = 300 + sum(row_heights)
    canvas = Image.new("RGB", (width * scale, height * scale), "white")
    draw = ImageDraw.Draw(canvas)
    navy, blue, gray = "#17324d", "#dbeaf5", "#aebdca"
    display_number = re.sub(r"(?<=\d)[ab]$", "", number)
    draw.text((46 * scale, 25 * scale), f"{display_number} {title}", font=font((28 if focused else 23) * scale), fill=navy)
    side_margin = 155 if focused else 125
    left, right = side_margin * scale, (width - side_margin) * scale
    xpos = {key: left + i * (right - left) // (len(lanes) - 1) for i, (key, _) in enumerate(lanes)}
    top = 88 * scale
    for key, label in lanes:
        x = xpos[key]
        box_half_width = min(140, (right - left) // ((len(lanes) - 1) * 2 * scale) - 8) if focused else 125
        long_label = len(label) > (16 if len(lanes) > 5 else 22)
        header_height = 64 if long_label else 48
        draw.rounded_rectangle((x - box_half_width * scale, top, x + box_half_width * scale, top + header_height * scale), radius=8 * scale, fill=blue, outline=navy, width=2 * scale)
        if long_label:
            display_label = re.sub(r"(Controller|Service|Repository|Storage|Mailer)$", r"\n\1", label)
            draw.multiline_text((x, top + 32 * scale), display_label, anchor="mm", align="center", spacing=2 * scale, font=font(19 * scale), fill=navy)
        else:
            draw.text((x, top + 24 * scale), label, anchor="mm", font=font((19 if len(lanes) > 5 else 24 if focused else 16) * scale), fill=navy)
        draw.line((x, top + header_height * scale, x, (height - 30) * scale), fill=gray, width=2 * scale)
    y_position = 260 * scale
    for i, (source, target, label, response) in enumerate(messages):
        if source not in xpos or target not in xpos:
            continue
        y = y_position
        a, b = xpos[source], xpos[target]
        color = "#597a91" if response else navy
        label_lines = labels[i]
        label_y = y - (len(label_lines) * 31 + 12) * scale
        if a == b:
            draw.arc((a, y - 13 * scale, a + 56 * scale, y + 45 * scale), 270, 90, fill=color, width=3 * scale)
            if focused and source == lanes[-2][0]:
                for j, part in enumerate(label_lines):
                    draw.text((a - 18 * scale, label_y + j * 31 * scale), part, anchor="ra", font=font(message_size * scale), fill=color)
            else:
                for j, part in enumerate(label_lines):
                    draw.text((a + 75 * scale, label_y + j * 31 * scale), part, font=font(message_size * scale), fill=color)
        else:
            draw.line((a, y, b, y), fill=color, width=3 * scale)
            direction = 1 if b > a else -1
            draw.polygon([(b, y), (b - direction * 13 * scale, y - 6 * scale), (b - direction * 13 * scale, y + 6 * scale)], fill=color)
            midpoint = (a + b) // 2
            for j, part in enumerate(label_lines):
                draw.text((midpoint, label_y + j * 31 * scale), part, anchor="ma", font=font(message_size * scale), fill=color)
        y_position += row_heights[i] * scale
    target_dir = output_dir or FIGURES
    target_dir.mkdir(parents=True, exist_ok=True)
    path = target_dir / (f"seq-{number[2:].lower()}.png" if output_dir else f"sequence-{number.lower()}.png")
    canvas.save(path)
    return path


def add_inline(paragraph, value: str) -> None:
    # Preserve link targets without placing long URLs in the report body.
    pieces = re.split(r"(\[[^]]+\]\([^)]+\))", value)
    for piece in pieces:
        match = re.fullmatch(r"\[([^]]+)\]\(([^)]+)\)", piece)
        if match:
            label, target = match.groups()
            if target.startswith("http"):
                link = OxmlElement("w:hyperlink")
                link.set(qn("r:id"), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
                run_xml = OxmlElement("w:r")
                props = OxmlElement("w:rPr")
                color = OxmlElement("w:color")
                color.set(qn("w:val"), "1765A0")
                props.append(color)
                run_xml.append(props)
                text_xml = OxmlElement("w:t")
                text_xml.text = label
                run_xml.append(text_xml)
                link.append(run_xml)
                paragraph._p.append(link)
            else:
                paragraph.add_run(label)
            continue
        for subpiece in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", piece):
            add_piece(paragraph, subpiece)


def add_piece(paragraph, piece: str) -> None:
    if not piece:
        return
    run = paragraph.add_run(piece.strip("`*"))
    run.bold = piece.startswith("**")
    fonts = run._element.get_or_add_rPr().rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.get_or_add_rPr().insert(0, fonts)
    fonts.set(qn("w:cs"), "Tahoma")
    fonts.set(qn("w:eastAsia"), "Tahoma")
    if piece.startswith("`"):
        run.font.name = "Consolas"
        run.font.size = Pt(10.5)


def add_bookmark(paragraph, name: str, number: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(number))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_page_ref(paragraph, name: str) -> None:
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    paragraph.add_run("\t")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), f"PAGEREF {name} \\h")
    run = OxmlElement("w:r")
    value = OxmlElement("w:t")
    value.text = "0"  # Updated by Word or LibreOffice when the document opens.
    run.append(value)
    field.append(run)
    paragraph._p.append(field)


def add_table(doc: Document, lines: list[str], caption: str | None = None,
              table_number: int = 0) -> None:
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)]
    if not rows:
        return
    if caption:
        p = doc.add_paragraph(caption, style="Caption")
        p.paragraph_format.keep_with_next = True
        add_bookmark(p, f"haris_tbl_{table_number:04d}", 3000 + table_number)
    cols = max(map(len, rows))
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    integrated_flow = cols == 3 and rows[0][:2] == ["รายการ", "รายละเอียด"] and any(row[0] == "Main Flow" for row in rows)
    query_table = rows[0] == ["Query ID", "UC ID", "UC name", "SQL command / source"]
    class_inventory = rows[0][:2] == ["Design class", "ชนิด"]
    table.autofit = not (integrated_flow or query_table or class_inventory)
    if integrated_flow:
        for column, width in zip(table.columns, (Cm(2.7), Cm(6.8), Cm(6.8))):
            column.width = width
        for row in table.rows:
            for cell, width in zip(row.cells, (Cm(2.7), Cm(6.8), Cm(6.8))):
                cell.width = width
    if query_table or class_inventory:
        widths = (1.4, 2.0, 3.0, 9.9) if query_table else (2.6, 3.0, 3.2, 7.5)
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    header_props = table.rows[0]._tr.get_or_add_trPr()
    header_props.append(OxmlElement("w:tblHeader"))
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            p = table.cell(i, j).paragraphs[0]
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, cell)
            for run in p.runs:
                run.font.size = Pt(10.5)
            if i == 0:
                for run in p.runs:
                    run.bold = True
        table.rows[i]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    if integrated_flow:
        flow_start = next(i for i, row in enumerate(rows) if row[0] == "Main Flow")
        flow_end = next(i for i in range(flow_start + 1, len(rows)) if rows[i][0]) - 1
        table.cell(flow_start, 0).merge(table.cell(flow_end, 0))
        for i, row in enumerate(rows):
            if i != flow_start and row[0]:
                table.cell(i, 1).merge(table.cell(i, 2))


def inventories(lines: list[str]) -> tuple[list[tuple[int, str]], list[str], list[str]]:
    """Collect ordered headings and captions before building the front matter."""
    headings: list[tuple[int, str]] = []
    figures: list[str] = []
    tables: list[str] = []
    chapter = 0
    figure_number = 0
    table_number = 0
    context = ""
    in_fence = False
    in_table = False
    for raw in lines:
        line = raw.strip()
        if line.startswith("```"):
            in_fence = not in_fence
            continue
        if in_fence:
            continue
        heading = re.match(r"^(#{2,5})\s+(.+)", line)
        if heading:
            level, label = len(heading.group(1)), heading.group(2)
            headings.append((level, label))
            context = label
            match = re.match(r"บทที่\s+(\d+)", label)
            if level == 2 and match:
                chapter = int(match.group(1))
                figure_number = table_number = 0
        image = re.fullmatch(r"!\[([^]]*)\]\([^)]+\)", line)
        if image:
            figure_number += 1
            label = re.sub(r"^ภาพ(?:ที่)?\s*\d+(?:\.\d+)*\s*", "", image.group(1)).removeprefix("ภาพ ")
            label = re.sub(r"^SEQ-(\d+[ab]?)\s+(.+)$",
                           lambda m: f"Sequence Diagram UC{m.group(1)}: {m.group(2)}", label)
            figures.append(f"ภาพที่ {chapter}.{figure_number} {label}")
        if line.startswith("|") and not in_table:
            table_number += 1
            label = re.sub(r"^\d+(?:\.\d+)*\s*", "", context)
            tables.append(f"ตารางที่ {chapter}.{table_number} {label}")
        in_table = line.startswith("|")
    return headings, figures, tables


def add_front_matter(doc: Document, headings: list[tuple[int, str]],
                     figures: list[str], tables: list[str]) -> None:
    for title, entries in (("สารบัญ", headings), ("สารบัญรูป", figures),
                           ("สารบัญตาราง", tables)):
        doc.add_page_break()
        doc.add_heading(title, 1)
        for number, entry in enumerate(entries, 1):
            if isinstance(entry, tuple):
                level, label = entry
                if level > 4:
                    continue
                p = doc.add_paragraph(label)
                p.paragraph_format.left_indent = Cm(max(0, level - 2) * 0.55)
                bookmark = f"haris_h_{number:04d}"
            else:
                p = doc.add_paragraph(entry)
                bookmark = f"haris_{'fig' if title == 'สารบัญรูป' else 'tbl'}_{number:04d}"
            p.paragraph_format.space_after = Pt(1)
            for run in p.runs:
                run.font.size = Pt(10.5)
            add_page_ref(p, bookmark)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    review_notes = next((i for i, line in enumerate(lines)
                         if line.startswith("### ประเด็นรอตรวจยืนยันก่อนจัดทำรายงานฉบับส่ง")),
                        len(lines))
    lines = lines[:review_notes]
    headings, figure_captions, table_captions = inventories(lines)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.2)
    styles = doc.styles
    font_sizes = {
        "Normal": 12.5,
        "Title": 24,
        "Heading 1": 20,
        "Heading 2": 17,
        "Heading 3": 15,
        "Heading 4": 13.5,
        "List Bullet": 12.5,
        "Caption": 11.5,
    }
    for name, size in font_sizes.items():
        style = styles[name]
        style.font.name = "Tahoma"
        fonts = style._element.get_or_add_rPr().rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            style._element.get_or_add_rPr().insert(0, fonts)
        fonts.set(qn("w:cs"), "Tahoma")
        fonts.set(qn("w:eastAsia"), "Tahoma")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(5)
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
    title_style = styles["Title"]
    title_ppr = title_style._element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)
    styles["Normal"].paragraph_format.line_spacing = 1.3
    for text, size, before, after in [
        ("รายงาน", 24, 36, 18),
        ("ระบบจัดการพนักงานและเงินเดือนร้าน Haris", 22, 0, 32),
        ("เสนอ", 16, 0, 12),
        ("[ชื่ออาจารย์ผู้สอน]", 16, 0, 32),
        ("คณะผู้จัดทำ", 18, 0, 12),
        ("[ชื่อและรหัสนักศึกษาของสมาชิก]", 16, 0, 36),
        ("รายงานนี้เป็นส่วนหนึ่งของรายวิชาการวิเคราะห์และออกแบบระบบ", 15, 0, 8),
        ("[รหัสวิชา ภาคเรียน และปีการศึกษา]", 15, 0, 8),
        ("[คณะ ภาควิชา และมหาวิทยาลัย]", 15, 0, 0),
    ]:
        paragraph = doc.add_paragraph(style="Title" if text == "รายงาน" else "Normal")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.space_after = Pt(after)
        add_inline(paragraph, text)
        for run in paragraph.runs:
            run.font.size = Pt(size)
    add_front_matter(doc, headings, figure_captions, table_captions)
    current_uc: tuple[str, str] | None = None
    heading_index = 0
    figure_index = 0
    table_index = 0
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            i += 1
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            heading = line[level:].strip()
            if level == 2 and doc.paragraphs:
                doc.add_page_break()
            heading_p = doc.add_heading(heading, min(level - 1, 4))
            # Keep the small bank-export example clear of the preceding full-page payslip image.
            if heading.startswith("3.10.1 CSV โอนเงิน"):
                heading_p.paragraph_format.page_break_before = True
            if level == 2 and heading.startswith("บทที่ "):
                heading_p.paragraph_format.space_before = Pt(48)
                heading_p.paragraph_format.space_after = Pt(28)
            if re.match(r"3\.4\.\d+ UC\d{2} ", heading):
                heading_p.paragraph_format.page_break_before = True
            heading_index += 1
            add_bookmark(heading_p, f"haris_h_{heading_index:04d}",
                         1000 + heading_index)
            match = re.match(r"3\.3\.\d+\s+Use Case ที่ \d+: (.+) \((UC\d+)\)", heading)
            if match:
                current_uc = (match.group(2), match.group(1))
            i += 1
            continue
        if line.startswith("|"):
            group = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                group.append(lines[i]); i += 1
            add_table(doc, group, table_captions[table_index], table_index + 1)
            table_index += 1
            continue
        image_match = re.fullmatch(r"!\[([^]]*)\]\(([^)]+)\)", line)
        if image_match:
            image_path = ROOT / image_match.group(2)
            if image_path.exists():
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
                with Image.open(image_path) as figure:
                    width_cm = min(15.5, 23.0 * figure.width / figure.height)
                paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))
                caption_p = doc.add_paragraph(figure_captions[figure_index], style="Caption")
                add_bookmark(caption_p, f"haris_fig_{figure_index + 1:04d}",
                             2000 + figure_index + 1)
                figure_index += 1
            i += 1
            continue
        if line.startswith("```mermaid"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i]); i += 1
            i += 1
            if block and block[0].strip() == "sequenceDiagram" and current_uc:
                image = draw_sequence(*current_uc, block)
                if image:
                    paragraph = doc.add_paragraph()
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.keep_with_next = True
                    paragraph.add_run().add_picture(str(image), width=Cm(15.5))
                    doc.add_paragraph(f"ภาพ Sequence Diagram {current_uc[0]}: {current_uc[1]}", style="Caption")
            else:
                p = doc.add_paragraph("แผนภาพเชิงแนวคิด: " + " / ".join(short(part.strip(), 70) for part in block if "-->" in part)[:450])
                p.style = "Caption"
            continue
        if line.startswith("```"):
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                p = doc.add_paragraph(lines[i])
                p.style = "No Spacing"
                for run in p.runs:
                    run.font.name = "Consolas"
                    run.font.size = Pt(10.5)
                i += 1
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:])
            i += 1
            continue
        p = doc.add_paragraph()
        if line.startswith("บทบาท ") and "ใช้หน้า" in line:
            p.paragraph_format.keep_with_next = True
        add_inline(p, line)
        i += 1
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("หน้า ")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    assert heading_index == len(headings), (heading_index, len(headings))
    assert figure_index == len(figure_captions), (figure_index, len(figure_captions))
    assert table_index == len(table_captions), (table_index, len(table_captions))
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)
    # Word uses a separate size property for Thai complex-script text.
    for root in (doc.element, doc.styles.element):
        for props in root.xpath(".//w:rPr"):
            size = props.find(qn("w:sz"))
            if size is not None:
                complex_size = props.find(qn("w:szCs"))
                if complex_size is None:
                    complex_size = OxmlElement("w:szCs")
                    props.append(complex_size)
                complex_size.set(qn("w:val"), size.get(qn("w:val")))
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
